"""按课程规划预留岗位；本脚本只生成本地预览，不连接飞书。"""
import argparse
import collections
import csv
import hashlib
import io
import json
import pathlib
import re

DEFAULT_HEADERS = json.loads((pathlib.Path(__file__).resolve().parent.parent / 'assets' / '单维表列序.json').read_text(encoding='utf-8'))['headers']
INITIAL_CHECK_ROLES = json.loads((pathlib.Path(__file__).resolve().parent.parent / 'assets' / '岗位适用项.json').read_text(encoding='utf-8'))['initial_check_roles']
ALIASES = {'课程名':['课程名','课程'], '群号':['群号','学习群号'],
           '职责':['职责','岗位'], '姓名':['姓名','助教姓名'], '微信号':['微信号','微信']}

def read_json(file):
    return json.loads(pathlib.Path(file).read_text(encoding='utf-8-sig'))

def letter(index):
    out = ''
    while index >= 0:
        index, digit = divmod(index, 26)
        out = chr(65 + digit) + out
        index -= 1
    return out

def columns(headers):
    found = {}
    for key, names in ALIASES.items():
        hits = [i for i, h in enumerate(headers) if h in names]
        if len(hits) != 1:
            raise ValueError('表头无法唯一定位：' + key)
        found[key] = hits[0]
    return found

def read_snapshot(file):
    data = read_json(file)
    if data.get('ok') is False:
        raise ValueError('飞书源表读取失败')
    data = data.get('data', data)
    if data.get('has_more') or data.get('truncated'):
        raise ValueError('源表未读完整，禁止生成写入计划')
    if 'rows' in data and 'headers' in data:
        return data['headers'], data['rows'], data.get('revision')
    text = data.get('annotated_csv', '')
    text = re.sub(r'^\[row=\d+\]\s*', '', text, flags=re.M)
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise ValueError('源表没有表头')
    return rows[0], rows[1:], data.get('revision')

def read_export(file, aliases=None):
    import openpyxl
    book = openpyxl.load_workbook(file, read_only=True, data_only=True)
    rows = book.active.iter_rows(values_only=True)
    headers = next(rows)
    required = ['任务名称','审核状态','表单数据']
    if any(k not in headers for k in required):
        raise ValueError('报名导出缺少必要表头：' + ','.join(required))
    people = []
    for number, values in enumerate(rows, 2):
        raw = dict(zip(headers, values))
        if not raw.get('任务名称'):
            continue
        title = str(raw['任务名称'])
        course, separator, role = title.rpartition(' - ')
        if not separator:
            course, role = title, ''
        status = str(raw.get('审核状态') or '')
        form = {}
        if status != '已释放':
            try:
                form = json.loads(raw.get('表单数据') or '{}')
            except (ValueError, TypeError):
                raise ValueError('第 %d 行报名表单不是有效 JSON' % number)
            if not isinstance(form, dict):
                raise ValueError('第 %d 行报名表单必须为对象' % number)
        pending = status == '已领取'
        people.append({'course':(aliases or {}).get(course, course),
                       'role':role.replace('助教',''),
                       'name':'' if pending else str(form.get('姓名') or ''),
                       'wechat':'' if pending else str(form.get('微信号（方便找到你）') or form.get('微信号') or ''),
                       'status':status, 'source_row':number})
    book.close()
    return people

def read_offline(file, period):
    """读取待同步 Excel 的身份列；其他人工过程列仍以线上现状为准。"""
    import openpyxl
    book = openpyxl.load_workbook(file, read_only=True, data_only=True)
    try:
        sheet = book.active
        if not sheet.title.startswith(period):
            raise ValueError('离线清单营期不一致')
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        col = columns(headers)
        people, course, group = [], '', ''
        for number, cells in enumerate(rows, 2):
            if cells[col['课程名']]:
                course = str(cells[col['课程名']]).strip()
            if cells[col['群号']]:
                group = str(cells[col['群号']]).strip()
            role = str(cells[col['职责']] or '').strip().replace('助教','')
            name = str(cells[col['姓名']] or '').strip()
            wechat = str(cells[col['微信号']] or '').strip()
            if not (name or wechat):
                continue
            if not course or not group or role not in ('运营','专业'):
                raise ValueError('离线清单身份列缺少课程、群号或岗位：第 %d 行' % number)
            people.append({'course':course,'role':role,'name':name,'wechat':wechat,
                           'status':'已提交','source_row':number})
        return people
    finally:
        book.close()

def build_plan(period, courses, people, headers=None, previous=None):
    headers = headers or DEFAULT_HEADERS.copy()
    previous = previous or []
    col = columns(headers)
    managed = set(col.values())
    conflicts, unallocated, warnings = [], [], []
    old_by_key = collections.defaultdict(list)
    old_rows = []
    course, group = '', ''
    for offset, cells in enumerate(previous, 2):
        row = list(cells) + [''] * max(0, len(headers)-len(cells))
        row = row[:len(headers)]
        if row[col['课程名']]:
            next_course = str(row[col['课程名']]).strip()
            if next_course != course:
                group = ''
            course = next_course
        group = str(row[col['群号']] or group).strip()
        role = str(row[col['职责']] or '').replace('助教','').strip()
        name, wechat = str(row[col['姓名']] or '').strip(), str(row[col['微信号']] or '').strip()
        if not any(str(v or '').strip() for v in row):
            continue
        item = {'row':offset,'cells':row,'course':course,'group':group,'role':role,'name':name,'wechat':wechat}
        old_rows.append(item)
        if name or wechat:
            old_by_key[(course,role,wechat or name)].append(item)
    known = {c['name'] for c in courses}
    if len(known) != len(courses):
        raise ValueError('课程名重复')
    members = collections.defaultdict(lambda: {'运营':[], '专业':[]})
    seen = set()
    released = 0
    for person in people:
        if person['status'] == '已释放':
            released += 1
            continue
        if person['course'] not in known or person['role'] not in ('运营','专业'):
            unallocated.append({'source_row':person['source_row'],'reason':'课程或岗位无法匹配'})
            continue
        key = (person['course'], person['role'], person['wechat'] or person['name'])
        if key[2] and key in seen:
            conflicts.append({'source_row':person['source_row'],'reason':'同课程同岗位人员标识重复，需核对报名'})
        seen.add(key)
        if person['status'] != '已领取' and (not person['name'] or not person['wechat']):
            warnings.append({'source_row':person['source_row'],'reason':'姓名或微信尚未补齐'})
        members[person['course']][person['role']].append(person)
    rows, merges, identities = [], [], []
    used_old = set()
    for c in courses:
        name, groups = c['name'], c['group_count']
        if isinstance(groups,bool) or not isinstance(groups,int) or groups <= 0:
            raise ValueError('群数必须为正整数：' + name)
        first = len(rows) + 2
        for role, capacity in [('运营',groups),('专业',groups*2)]:
            for p in members[name][role][capacity:]:
                unallocated.append({'source_row':p['source_row'],'reason':'报名超过规划岗位数'})
        for g in range(groups):
            group_first = len(rows)+2
            for slot, role in enumerate(['运营','专业','专业']):
                index = g if slot == 0 else g*2+slot-1
                candidates = members[name][role]
                person = candidates[index] if index < len(candidates) else None
                row = ['']*len(headers)
                if person and (person['wechat'] or person['name']):
                    key = (name,role,person['wechat'] or person['name'])
                    hits = old_by_key.get(key,[])
                    if len(hits)>1:
                        conflicts.append({'reason':'源表同课程同岗位人员匹配不唯一','source_rows':[h['row'] for h in hits]})
                    elif hits:
                        old = hits[0]
                        used_old.add(old['row'])
                        row = old['cells'].copy()
                        if old['group'] and old['group'] != '%d群'%(g+1):
                            conflicts.append({'reason':'平台顺序将改变已有人员群号，先核对人工调整','source_row':old['row']})
                if not any(str(v or '').strip() for v in row):
                    # 空岗位也可有人工观察员等信息；原位置身份为空时保留。
                    destination = len(rows)+2
                    empty = next((o for o in old_rows if o['row']==destination and not o['name'] and not o['wechat']
                                  and o['course']==name and o['role']==role and o['group']=='%d群'%(g+1)),None)
                    if empty:
                        row = empty['cells'].copy(); used_old.add(empty['row'])
                row[col['课程名']] = name if g==0 and slot==0 else ''
                row[col['群号']] = '%d群'%(g+1) if slot==0 else ''
                row[col['职责']] = role
                row[col['姓名']] = person['name'] if person else ''
                row[col['微信号']] = person['wechat'] if person else ''
                # 只补空检查项；按表头定位，保留人工状态、评分和其它信息。
                for header, applicable_roles in INITIAL_CHECK_ROLES.items():
                    if header in headers:
                        check_index = headers.index(header)
                        if row[check_index] is None or row[check_index] == '':
                            row[check_index] = '待核对' if role in applicable_roles else '不适用'
                if '备注' in headers:
                    note_index=headers.index('备注')
                    marker='平台状态：已领取未提交'
                    lines=[x for x in str(row[note_index] or '').splitlines() if x.strip()!=marker]
                    if person and person['status']=='已领取':
                        lines.append(marker)
                    row[note_index]='\n'.join(lines)
                rows.append(row)
                identities.append({'business_key':json.dumps([period,name,g+1,role,slot],ensure_ascii=False),
                                   'source_row':person['source_row'] if person else None,
                                   'state':person['status'] if person else '空缺'})
            merges.append('%s%d:%s%d'%(letter(col['群号']),group_first,letter(col['群号']),group_first+2))
        merges.insert(len(merges)-groups,'%s%d:%s%d'%(letter(col['课程名']),first,letter(col['课程名']),len(rows)+1))
    for old in old_rows:
        if old['row'] not in used_old and any(str(v or '').strip() for i,v in enumerate(old['cells']) if i not in managed):
            conflicts.append({'source_row':old['row'],'reason':'未匹配旧记录含人工过程数据，禁止覆盖或清除'})
    diff=[]
    for n,row in enumerate(rows):
        prior = previous[n] if n<len(previous) else []
        for i,value in enumerate(row):
            before = prior[i] if i<len(prior) else ''
            if str(before or '') != str(value or ''):
                diff.append({'cell':letter(i)+str(n+2),'before':before,'after':value})
    if len(previous)>len(rows) and any(any(str(v or '').strip() for v in r) for r in previous[len(rows):]):
        conflicts.append({'reason':'目标表尾部还有旧记录，先明确处理范围，不自动删除'})
    return {'period':period,'headers':headers,'rows':rows,'identities':identities,
            'merge_ranges':merges,'diff':diff,'unallocated':unallocated,'conflicts':conflicts,'warnings':warnings,
            'write_ready':not conflicts and not unallocated,
            'summary':{'courses':len(courses),'groups':sum(c['group_count'] for c in courses),
                       'slots':len(rows),'retained':sum(p['status']!='已释放' for p in people),
                       'released':released,'vacancies':sum(i['state']=='空缺' for i in identities),
                       'changed_cells':len(diff)}}

def write_workbook(plan, file):
    """将已生成的名单计划保存为可编辑的离线清单，不代表线上已同步。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    book = Workbook()
    sheet = book.active
    sheet.title = plan['period'] + '待同步'
    sheet.append(plan['headers'])
    for row in plan['rows']:
        sheet.append(row)
    sheet.freeze_panes = 'F2'
    colors = {'运营':'DBEAFE', '专业':'DCFCE7'}
    inactive_color = 'E5E7EB'
    role_index = plan['headers'].index('职责') + 1
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E5E7EB')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        role = row[role_index-1].value
        if role in colors:
            row[role_index-1].fill = PatternFill('solid', fgColor=colors[role])
        for header, roles in INITIAL_CHECK_ROLES.items():
            if header in plan['headers']:
                cell = row[plan['headers'].index(header)]
                cell.fill = PatternFill('solid', fgColor=colors[role] if role in roles else inactive_color)
        for cell in row:
            cell.protection = Protection(locked=False)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    for header in INITIAL_CHECK_ROLES:
        if header in plan['headers']:
            index = plan['headers'].index(header) + 1
            validation = DataValidation(type='list',
                                        formula1='"待核对,已核实完成,存在问题,不适用"',
                                        allow_blank=False)
            sheet.add_data_validation(validation)
            column = get_column_letter(index)
            validation.add(f'{column}2:{column}{max(2, len(plan["rows"])+1)}')
    for index, header in enumerate(plan['headers'], 1):
        sheet.column_dimensions[get_column_letter(index)].width = 22 if header in ('课程名','微信号') else 16
    sheet.row_dimensions[1].height = 30
    for merged in plan['merge_ranges']:
        sheet.merge_cells(merged)
    book.save(file)
    book.close()

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--period',required=True)
    parser.add_argument('--courses',required=True)
    parser.add_argument('--export',required=True)
    parser.add_argument('--offline')
    parser.add_argument('--target')
    parser.add_argument('--aliases')
    parser.add_argument('--out',required=True)
    args=parser.parse_args()
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',args.period):
        raise ValueError('营期必须为 YYYY-MM')
    cache=read_json(args.courses); courses=cache.get('courses',[])
    aliases=read_json(args.aliases) if args.aliases else {}
    headers, previous, revision=read_snapshot(args.target) if args.target else (None,None,None)
    people=read_offline(args.offline,args.period) if args.offline else read_export(args.export,aliases)
    plan=build_plan(args.period,courses,people,headers,previous)
    plan['target_revision']=revision
    plan['input_hashes']={label:hashlib.sha256(pathlib.Path(file).read_bytes()).hexdigest()
                          for label,file in [('courses',args.courses),('offline' if args.offline else 'export',args.offline or args.export)]+([('target',args.target)] if args.target else [])}
    out=pathlib.Path(args.out)
    out.mkdir(parents=True,exist_ok=False)
    (out/'名单写入计划.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    lines=['# 助教清单预览','',f"课程 {plan['summary']['courses']}，群数 {plan['summary']['groups']}，岗位 {len(plan['rows'])}，空缺 {plan['summary']['vacancies']}。",
           f"排除已释放 {plan['summary']['released']}；冲突 {len(plan['conflicts'])}；未分配 {len(plan['unallocated'])}。",
           '', '详细姓名、微信及单元格差异保存在同目录 JSON，仅供本地核验。']
    (out/'名单预览.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    write_workbook(plan,out/'助教清单_待同步.xlsx')
    print(json.dumps({'write_ready':plan['write_ready'],**plan['summary']},ensure_ascii=False))

if __name__=='__main__':
    main()
